"""Streamlit interface for Cu(II)-microplastic sorption prediction and interpretation (revised model).

Run with:
    streamlit run src/ui/app.py

The application uses the model selected in the revision (pooled Freundlich isotherm with polymer, ageing
status and adsorption system; ``src/analysis/predictor.py``) and the LLM Interpretation System of
``src/rag/llm_interpretation_v2.py``. The interface of the submitted version (Quantile Random Forest)
is kept in ``src/ui/app_submitted_version.py``.
"""

from __future__ import annotations

import html
import io
import re
import sys
import threading
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter, NullFormatter
import pandas as pd
import streamlit as st

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAG_DIR = PROJECT_ROOT / "src" / "rag"
for path in (PROJECT_ROOT, RAG_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import config  # noqa: E402
from llm_clients import GENERATOR, available  # noqa: E402
from llm_interpretation_v2 import LABELS, explain, graph_evidence, readable_text, retrieval_available  # noqa: E402
from src.analysis.predictor import MIN_STUDIES, fmt_qe, fmt_range, model_summary, predict, sig  # noqa: E402

st.set_page_config(page_title="Cu²⁺-MP Predictor", layout="wide", page_icon="\U0001f9ea", initial_sidebar_state="expanded")

st.markdown("""
<style>
    [data-testid="metric-container"] { background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 12px 16px; }
    [data-testid="stTab"] { font-weight: 600; }
    .section-header { border-left: 4px solid #4c78a8; padding-left: 10px; margin-bottom: 8px; }
    .badge-ok   { color: #16803c; font-weight: 700; }
    .badge-fail { color: #b42318; font-weight: 700; }
</style>
""", unsafe_allow_html=True)

INPUT_COLS = ["ReT", "AgS", "AdC", "Ce"]
PARAM_LABELS = {"ReT": "Resin type (polymer)", "AgS": "Ageing status", "AdC": "Adsorption condition",
                "Ce": "Equilibrium Cu²⁺ concentration, Ce (mg/L)"}
MODE_OPTIONS = ["C1 — Rule-based (no LLM)", "C2 — Retrieval", "C3 — Retrieval + knowledge graph"]
REPOSITORY = "https://github.com/nguyencuongqt/cu-mp-adsorption-ml"
CORPUS_LIST = config.MODELS_DIR / "corpus_v2_papers.csv"
MAX_BATCH_EXPLANATIONS = 50


# ── Helpers without interface elements ───────────────────────────────────────

def _missing(x: Any) -> bool:
    return x is None or (isinstance(x, str) and not x.strip()) or (not isinstance(x, str) and pd.isna(x))


def anchor_of(values: dict[str, Any]) -> tuple[float, float] | None:
    """(Ce, qe) of the local measurement, if one was entered. An incomplete or non-positive pair raises an error
    instead of being ignored, so that the user knows the narrower interval was not applied."""
    ce, qe = values.get("anchor_Ce"), values.get("anchor_qe")
    if _missing(ce) and _missing(qe):
        return None
    if _missing(ce) or _missing(qe):
        raise ValueError("The local measurement needs both Ce (mg/L) and qe (mg/g).")
    try:
        ce, qe = float(ce), float(qe)
    except (TypeError, ValueError):
        raise ValueError("The local measurement must be numeric: Ce in mg/L and qe in mg/g.") from None
    if not (np.isfinite(ce) and np.isfinite(qe) and ce > 0 and qe > 0):
        raise ValueError("The local measurement needs a positive Ce (mg/L) and a positive qe (mg/g).")
    return ce, qe


def run_prediction(values: dict[str, Any]) -> dict[str, Any]:
    return predict({k: values[k] for k in INPUT_COLS}, anchor=anchor_of(values))


def batch_predict(df: pd.DataFrame, defaults: dict[str, Any]) -> pd.DataFrame:
    """One prediction per row; columns missing from the file take the sidebar values."""
    rows = []
    for _, row in df.iterrows():
        raw = {k: row[k] if k in row.index and pd.notna(row[k]) else defaults.get(k) for k in INPUT_COLS}
        for k in ("anchor_Ce", "anchor_qe"):
            raw[k] = row[k] if k in row.index else None
        try:
            pred = run_prediction(raw)
            rows.append({**pred["inputs"], "anchor_Ce": raw["anchor_Ce"], "anchor_qe": raw["anchor_qe"], "lab_qe": pred["lab_qe"], "lab_qe_lower": pred["lab_qe_lower"],
                         "lab_qe_upper": pred["lab_qe_upper"], "interval_factor": pred["interval_factor"],
                         "local_measurement_used": pred["anchored"], "river_qe": pred["river_qe"],
                         "river_qe_lower": pred["river_qe_lower"], "river_qe_upper": pred["river_qe_upper"],
                         "river_interval_factor": pred["river_interval_factor"],
                         **{f"factor_{d['predictor']}": d["factor_on_qe"] for d in pred["drivers"]},
                         "warnings": " | ".join(pred["warnings"])})
        except Exception as exc:  # invalid row: keep it in the output with the reason
            rows.append({k: raw.get(k) for k in INPUT_COLS} | {"error": str(exc)})
    return pd.DataFrame(rows)


def isotherm_figure(pred: dict[str, Any]) -> plt.Figure:
    """Predicted isotherm with its 90% interval, the laboratory concentration range and the requested point."""
    s, v = model_summary(), pred["inputs"]
    ce = np.logspace(np.log10(min(1e-4, v["Ce"] / 3)), np.log10(max(1e3, v["Ce"] * 3)), 200)
    qe = pred["lab_qe"] * (ce / v["Ce"]) ** pred["freundlich_exponent"]
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    ax.axvspan(s["ce_min"], s["ce_max"], color="#e8eef5", label="Laboratory Ce range")
    ax.axvspan(s["ce_min"], s["ce_dense"], color="#f6e7c8", label=f"Sparse laboratory data (lowest 5%)")
    ax.fill_between(ce, qe / pred["interval_factor"], qe * pred["interval_factor"], color="#4c78a8", alpha=0.22,
                    label=f"90% prediction interval (×/÷{sig(pred['interval_factor'])})")
    ax.plot(ce, qe, color="#4c78a8", lw=2, label="Laboratory-equilibrium qe")
    ax.plot(ce, qe * pred["field_ratio"], color="#e45756", lw=1.6, ls="--", label=f"Field-scale value (ratio {sig(pred['field_ratio'])})")
    ax.scatter([v["Ce"]], [pred["lab_qe"]], color="#1f3b57", zorder=5)
    ax.scatter([v["Ce"]], [pred["river_qe"]], color="#e45756", zorder=5)
    ax.set_xscale("log")
    ax.set_yscale("log")
    # Plain tick labels (no mathtext): the mathtext parser can fail when Streamlit interrupts a rerun while it runs.
    plain = FuncFormatter(lambda v, _: f"{v:g}")
    for axis in (ax.xaxis, ax.yaxis):
        axis.set_major_formatter(plain)
        axis.set_minor_formatter(NullFormatter())
    ax.set_xlabel("Ce (mg/L)")
    ax.set_ylabel("qe (mg/g)")
    ax.set_title(f"{v['ReT']}, {v['AgS'].lower()}, {v['AdC'].lower()} system")
    ax.legend(fontsize=8, loc="lower right", frameon=True, framealpha=0.9)
    ax.grid(alpha=0.25, which="major")
    fig.tight_layout()
    return fig


def drivers_table(pred: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{"Predictor": LABELS[d["predictor"]], "Value": str(pred["inputs"][d["predictor"]]),
                          "Multiplier of qe": sig(d["factor_on_qe"]), "Effect": d["direction"]} for d in pred["drivers"]])


def drivers_figure(pred: dict[str, Any]) -> plt.Figure:
    """Contribution of each input, as a multiplier of qe relative to the average laboratory observation."""
    drivers = pred["drivers"][::-1]
    effects = [d["log10_effect"] for d in drivers]
    fig, ax = plt.subplots(figsize=(7.2, 2.9))
    bars = ax.barh([LABELS[d["predictor"]] for d in drivers], effects, color=["#54a24b" if e > 0 else "#e45756" for e in effects])
    for bar, d in zip(bars, drivers):
        ax.text(bar.get_width(), bar.get_y() + bar.get_height() / 2, f"  ×{sig(d['factor_on_qe'])}  ",
                va="center", ha="left" if d["log10_effect"] > 0 else "right", fontsize=9)
    ax.axvline(0, color="#444", lw=0.8)
    ax.set_xlabel("Effect on log₁₀ qe relative to the average laboratory observation")
    ax.margins(x=0.25)
    fig.tight_layout()
    return fig


# Matplotlib is not thread-safe and Streamlit runs each session in its own thread; concurrent reruns can fail while the
# tick labels are parsed. Every figure is therefore drawn and saved under one lock and shown as an image.
@st.cache_resource
def figure_lock() -> threading.RLock:
    """One lock for all sessions (the script is re-executed on every rerun, so a module-level lock would not be shared)."""
    return threading.RLock()


def render_figure(make, pred: dict[str, Any]) -> tuple[bytes, bytes]:
    """PNG for the screen and a 300-dpi PNG for download."""
    with figure_lock():
        fig = make(pred)
        try:
            out = []
            for dpi in (150, 300):
                buf = io.BytesIO()
                fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
                out.append(buf.getvalue())
        finally:
            plt.close(fig)
    return out[0], out[1]


def safe_html(text: str) -> str:
    """Explanation HTML without scripts, styles, event handlers or links."""
    text = re.sub(r"<(script|style|iframe)[\s\S]*?</\1>", "", str(text), flags=re.I)
    text = re.sub(r"\son\w+\s*=\s*(\"[^\"]*\"|'[^']*'|\S+)", "", text, flags=re.I)
    return re.sub(r"</?(a|img|iframe|script|style|link|meta|form|input)\b[^>]*>", "", text, flags=re.I)


def source_labels(sources: list[dict[str, Any]]) -> list[str]:
    """Articles used as context, with their DOI."""
    doi = pd.read_csv(CORPUS_LIST).set_index("key")["doi"].to_dict() if CORPUS_LIST.exists() else {}
    labels = []
    for key in dict.fromkeys(str(s.get("paper_id")) for s in sources):
        d = doi.get(key)
        labels.append(f"{key} — https://doi.org/{d}" if isinstance(d, str) and d else key)
    return labels


MODE_NAMES = {0: "numbers only", 1: "C1 rule-based, no LLM", 2: "C2 retrieval, LLM-written",
              3: "C3 retrieval and knowledge graph, LLM-written"}


def report_text(result: dict[str, Any]) -> str:
    """Plain-text report: model output and explanation."""
    mode = MODE_NAMES[result["mode"]]
    lines = ["Cu(II)-microplastic sorption prediction", "=" * 40, "", "MODEL OUTPUT", result["facts"], "",
             f"EXPLANATION ({mode})", readable_text(result["explanation_html"])]
    sources = source_labels(result.get("sources", []))
    if sources:
        lines += ["", "SOURCES USED AS CONTEXT"] + [f"- {s}" for s in sources]
    return "\n".join(lines) + "\n"


def report_docx(result: dict[str, Any]) -> bytes:
    """The same report as a Word document."""
    from docx import Document

    doc = Document()
    doc.add_heading("Cu(II)-microplastic sorption prediction", level=1)
    doc.add_heading("Model output", level=2)
    for line in result["facts"].split("\n"):
        doc.add_paragraph(line)
    doc.add_heading(f"Explanation ({MODE_NAMES[result['mode']]})", level=2)
    for line in readable_text(result["explanation_html"]).split("\n"):
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    sources = source_labels(result.get("sources", []))
    if sources:
        doc.add_heading("Sources used as context", level=2)
        for s in sources:
            doc.add_paragraph(s, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def get_explanation(values: dict[str, Any], mode: int) -> dict[str, Any]:
    """Explanation for the case; results are kept for the session so that a case is sent to the LLM once."""
    cache = st.session_state.setdefault("explanations", {})
    key = (tuple(str(values.get(k)) for k in INPUT_COLS + ["anchor_Ce", "anchor_qe"]), mode)
    if key not in cache:
        cache[key] = explain({k: values[k] for k in INPUT_COLS}, mode=mode, anchor=anchor_of(values))
    return cache[key]


# ── Sidebar ──────────────────────────────────────────────────────────────────

def status_badge(label: str, ok: bool) -> None:
    css, text = ("badge-ok", "OK") if ok else ("badge-fail", "Missing")
    st.sidebar.markdown(f"{html.escape(label)}: <span class='{css}'>{text}</span>", unsafe_allow_html=True)


def render_sidebar() -> tuple[dict[str, Any], bool]:
    s = model_summary()
    st.sidebar.title("\U0001f9ea Cu²⁺-MP Predictor")
    st.sidebar.caption("Pooled Freundlich model validated by study + LLM interpretation system")
    st.sidebar.divider()

    with st.sidebar.expander("\U0001f4cb Model & component status", expanded=False):
        st.caption(f"{s['n_points']} observations, {s['n_studies']} laboratory studies; Freundlich exponent {s['exponent']:.2f}")
        status_badge("Retrieval index (C2, C3)", retrieval_available())
        status_badge("Corpus list", CORPUS_LIST.exists())
        status_badge(f"LLM key ({GENERATOR[1]})", available(GENERATOR[0]))

    st.sidebar.subheader("Model inputs")
    counts = s["studies_per_polymer"]
    polymers = sorted(s["levels"]["ReT"], key=lambda p: -counts[p])
    ret_val = st.sidebar.selectbox(PARAM_LABELS["ReT"], polymers, index=0,
                                   format_func=lambda p: f"{p} ({counts[p]} stud{'y' if counts[p] == 1 else 'ies'})",
                                   help=f"Polymers with fewer than {MIN_STUDIES} laboratory studies are flagged in the output.")
    ags_val = st.sidebar.radio(PARAM_LABELS["AgS"], s["levels"]["AgS"], horizontal=True,
                               help="Virgin = as produced; aged = weathered in the environment or in the laboratory.")
    adc_val = st.sidebar.radio(PARAM_LABELS["AdC"], s["levels"]["AdC"], horizontal=True,
                               help="Single-ion: Cu²⁺ only. Mixed: other metal ions compete.")
    ce_val = st.sidebar.number_input(PARAM_LABELS["Ce"], min_value=0.00001, max_value=1000.0, value=0.00148, format="%.5f",
                                     help=f"Laboratory range: {s['ce_min']:.2g}–{s['ce_max']:.3g} mg/L. "
                                          "World-average dissolved Cu in rivers: 0.00148 mg/L.")
    if ce_val < s["ce_min"] or ce_val > s["ce_max"]:
        st.sidebar.markdown("<span class='badge-fail'>Outside the laboratory range: extrapolation</span>", unsafe_allow_html=True)
    elif ce_val < s["ce_dense"]:
        st.sidebar.markdown("<span style='color:#b45309;font-weight:700'>Sparse laboratory data at this Ce</span>", unsafe_allow_html=True)
    else:
        st.sidebar.markdown("<span class='badge-ok'>Within the laboratory range</span>", unsafe_allow_html=True)
    st.sidebar.caption("Temperature, pH and shaking speed are not predictors of the revised model.")

    st.sidebar.divider()
    st.sidebar.subheader("Local measurement (optional)")
    values: dict[str, Any] = {"ReT": ret_val, "AgS": ags_val, "AdC": adc_val, "Ce": float(ce_val), "anchor_Ce": None, "anchor_qe": None}
    if st.sidebar.toggle("One measured (Ce, qe) pair for this material", value=False,
                         help=f"Shifts the isotherm through the measurement and narrows the 90% interval from "
                              f"×/÷{sig(s['factor_new_lab'])} to ×/÷{sig(s['factor_anchored'])}."):
        values["anchor_Ce"] = st.sidebar.number_input("Measured Ce (mg/L)", min_value=0.00001, value=1.0, format="%.5f")
        values["anchor_qe"] = st.sidebar.number_input("Measured qe (mg/g)", min_value=0.000001, value=0.1, format="%.6f")

    st.sidebar.divider()
    st.sidebar.subheader("Field-scale value")
    st.sidebar.caption(f"Field-to-laboratory ratio **{sig(s['field_ratio'])}** (95% CI {sig(s['field_ratio_low'])}–"
                       f"{sig(s['field_ratio_high'])}; {s['n_field_studies']} field studies). River qe = laboratory qe × ratio; "
                       "its interval also includes the variation of the ratio between rivers.")
    st.sidebar.divider()
    batch = st.sidebar.toggle("Batch mode (CSV/XLSX upload)", value=False)
    st.sidebar.divider()
    with st.sidebar.expander("ℹ️ About", expanded=False):
        st.markdown(f"Model, validation and field-to-laboratory ratio: manuscript *From laboratory isotherms to rivers: machine "
                    f"learning prediction of copper adsorption on microplastics, field scaling and LLM-assisted explanation*. "
                    f"Code and data: [{REPOSITORY.split('github.com/')[1]}]({REPOSITORY}). Predictions are screening estimates; "
                    "the intervals, not the central values, describe what is known.")
    return values, batch


# ── Panel A – input summary ──────────────────────────────────────────────────

def panel_input_summary(values: dict[str, Any], batch: bool) -> None:
    batch_out: pd.DataFrame | None = st.session_state.get("batch_results")
    if batch and batch_out is not None and not batch_out.empty:
        ok = batch_out[batch_out["lab_qe"].notna()] if "lab_qe" in batch_out.columns else batch_out.iloc[0:0]
        st.success(f"Batch mode — {len(batch_out)} rows, {len(ok)} predicted")
        if not ok.empty:
            col1, col2 = st.columns(2)
            col1.dataframe(ok["ReT"].value_counts().rename_axis("Polymer").reset_index(name="Rows"), hide_index=True, width="stretch")
            m = col2.columns(3)
            m[0].metric("Lowest qe", fmt_qe(ok['lab_qe'].min()))
            m[1].metric("Median qe", fmt_qe(ok['lab_qe'].median()))
            m[2].metric("Highest qe", fmt_qe(ok['lab_qe'].max()))
    elif batch:
        st.info("Batch mode — upload a file in panel B and run the prediction to see a summary here.")
    else:
        table = [(PARAM_LABELS[k], f"{values[k]:.5g}" if k == "Ce" else values[k]) for k in INPUT_COLS]
        if anchor_of(values):
            table.append(("Local measurement", f"Ce = {values['anchor_Ce']:.5g} mg/L, qe = {values['anchor_qe']:.5g} mg/g"))
        st.dataframe(pd.DataFrame(table, columns=["Input", "Value"]), hide_index=True, width="stretch")


# ── Panel B – prediction ─────────────────────────────────────────────────────

def show_warnings(pred: dict[str, Any]) -> None:
    for warning in pred["warnings"]:
        st.warning(f"⚠️ {warning}")
    if not pred["warnings"]:
        st.success("No validity warning applies to this input.")


def panel_prediction(values: dict[str, Any], batch: bool) -> None:
    if batch:
        upload = st.file_uploader("Upload CSV/XLSX with the columns " + ", ".join(INPUT_COLS)
                                  + " (optional: anchor_Ce, anchor_qe for one local measurement per row)", type=["csv", "xlsx"])
        if upload is not None and st.button("\U0001f9ee Run batch prediction", type="primary"):
            try:
                df = pd.read_csv(upload) if upload.name.lower().endswith(".csv") else pd.read_excel(upload)
            except Exception as exc:
                st.error(f"Could not read the file: {exc}")
                return
            st.session_state["batch_results"] = batch_predict(df, values)
            st.session_state.pop("batch_report", None)
            st.session_state.pop("explanation", None)
        batch_out: pd.DataFrame | None = st.session_state.get("batch_results")
        if batch_out is not None and not batch_out.empty:
            if "error" in batch_out.columns and batch_out["error"].notna().any():
                st.warning(f"{int(batch_out['error'].notna().sum())} row(s) could not be predicted; see the column 'error'.")
            st.dataframe(batch_out, width="stretch")
            st.download_button("\U0001f4be Download batch results (CSV)", batch_out.to_csv(index=False), "batch_predictions.csv", "text/csv")
        return

    if st.button("\U0001f9ee Run prediction", type="primary"):
        try:
            st.session_state["prediction"] = run_prediction(values)
            st.session_state["values"] = dict(values)
            st.session_state.pop("explanation", None)      # an explanation of the previous inputs must not remain visible
        except Exception as exc:
            st.error(str(exc))
            return
    pred = st.session_state.get("prediction")
    if not pred:
        st.info("ℹ️ Set the inputs in the sidebar and click **Run prediction**.")
        return

    if st.session_state.get("values") != values:
        st.warning("⚠️ The inputs in the sidebar have changed; click **Run prediction** to update the result.")
    left, right = st.columns(2)
    left.metric("Laboratory-equilibrium qe", fmt_qe(pred["lab_qe"]))
    left.caption(f"90% prediction interval: {fmt_range(pred['lab_qe_lower'], pred['lab_qe_upper'], pred['lab_qe'])} "
                 f"(×/÷{sig(pred['interval_factor'])}; " + ("one local measurement used)" if pred["anchored"] else "new laboratory)"))
    right.metric("Field-scale qe", fmt_qe(pred["river_qe"]))
    right.caption(f"90% interval for a river not among the field studies: "
                  f"{fmt_range(pred['river_qe_lower'], pred['river_qe_upper'], pred['river_qe'])} (×/÷{sig(pred['river_interval_factor'])}; "
                  f"model interval combined with the variation of the field-to-laboratory ratio between rivers)")
    show_warnings(pred)
    if not pred["anchored"]:
        st.caption(f"One measured (Ce, qe) pair for this material would narrow the laboratory interval to "
                   f"×/÷{sig(pred['interval_factor_with_one_local_measurement'])} (sidebar: Local measurement).")
    screen, hires = render_figure(isotherm_figure, pred)
    st.image(screen, width="stretch")
    st.download_button("\U0001f4be Download figure", hires, "prediction_isotherm.png", "image/png")


# ── Panel C – contributions of the inputs ────────────────────────────────────

def panel_drivers() -> None:
    pred = st.session_state.get("prediction")
    batch_out: pd.DataFrame | None = st.session_state.get("batch_results")
    if not pred:
        if batch_out is not None and any(c.startswith("factor_") for c in batch_out.columns):
            st.markdown("**Multiplier of qe from each input, by row** (relative to the average laboratory observation)")
            cols = [c for c in INPUT_COLS if c in batch_out.columns] + [c for c in batch_out.columns if c.startswith("factor_")]
            st.dataframe(batch_out[cols], width="stretch")
        else:
            st.info("ℹ️ Run a prediction first (panel B).")
        return
    st.markdown("Contribution of each input to this prediction, as a multiplier of qe relative to the average laboratory "
                "observation. The model is additive on the log scale, so the contributions are exact.")
    screen, hires = render_figure(drivers_figure, pred)
    st.image(screen, width="stretch")
    st.dataframe(drivers_table(pred), hide_index=True, width="stretch")
    st.download_button("\U0001f4be Download figure", hires, "input_contributions.png", "image/png")
    st.caption("SHAP values of the machine-learning benchmark models and their stability across studies are reported in the manuscript.")


# ── Panel D – LLM interpretation ─────────────────────────────────────────────

def show_explanation(result: dict[str, Any], key: str) -> None:
    st.html(safe_html(result["explanation_html"]))
    if result["generation"].get("attempts", 1) > 1:
        st.caption("The explanation was generated a second time because the first version contained a value that did not match the model output.")
    if result["generation"].get("unsupported_after_check", 0):
        st.warning("⚠️ The explanation contains a value in mg/g that does not match the model output; rely on the numbers in panel B.")
    if result["mode"] >= 2:
        labels = source_labels(result["sources"])
        with st.expander(f"\U0001f4da Sources used as context ({len(labels)} articles)"):
            for label in labels:
                st.markdown(f"- {label}")
    if result["mode"] == 3:
        v = result["prediction"]["inputs"]
        with st.expander("\U0001f578️ Knowledge-graph relations for this polymer"):
            st.text(graph_evidence(v["ReT"], v["AgS"])[0])
    c1, c2, c3 = st.columns(3)
    c1.download_button("\U0001f4be Report (.txt)", report_text(result), "cu_mp_report.txt", "text/plain", key=f"{key}_txt")
    try:
        c2.download_button("\U0001f4be Report (.docx)", report_docx(result), "cu_mp_report.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"{key}_docx")
    except ImportError:
        c2.caption("Install python-docx to export .docx")
    c3.download_button("\U0001f4be Explanation (.html)", safe_html(result["explanation_html"]), "explanation.html", "text/html", key=f"{key}_html")


def row_values(batch_ok: pd.DataFrame, i: int) -> dict[str, Any]:
    """Inputs of one batch row, including its local measurement if the file gave one."""
    row = batch_ok.loc[i]
    return {k: row[k] for k in INPUT_COLS} | {k: (row[k] if k in row.index and not _missing(row[k]) else None)
                                               for k in ("anchor_Ce", "anchor_qe")}


def panel_llm() -> None:
    values = st.session_state.get("values") if st.session_state.get("prediction") else None
    batch_ok: pd.DataFrame | None = None
    if values is None:
        batch_out: pd.DataFrame | None = st.session_state.get("batch_results")
        if batch_out is None or "lab_qe" not in batch_out.columns or batch_out["lab_qe"].notna().sum() == 0:
            st.info("ℹ️ Run a prediction first (panel B).")
            return
        batch_ok = batch_out[batch_out["lab_qe"].notna()].reset_index(drop=True)
        idx = st.selectbox("Row to explain", range(len(batch_ok)), key="llm_batch_row",
                           format_func=lambda i: f"Row {i + 1} | {batch_ok.loc[i, 'ReT']} | {batch_ok.loc[i, 'AgS']} | Ce = {batch_ok.loc[i, 'Ce']:.4g} mg/L")
        values = row_values(batch_ok, idx)

    api_ok, index_ok = available(GENERATOR[0]), retrieval_available()
    mode = int(st.radio("Interpretation mode", MODE_OPTIONS, horizontal=True, key="llm_mode")[1])
    if mode >= 2 and not index_ok:
        st.info("ℹ️ C2 and C3 need the retrieval index built from the 41 source articles, which cannot be redistributed. "
                f"They run in a local installation (see the README of the [repository]({REPOSITORY})); the C2 and C3 explanations "
                "evaluated in the manuscript can be read in panel E. C1 is used here.")
        mode = 1
    elif mode >= 2 and not api_ok:
        st.warning("⚠️ C2 and C3 need an API key for the LLM service (OPENAI_API_KEY in .env or in the app secrets). C1 is used instead.")
        mode = 1
    if mode >= 2:
        st.caption(f"C2 and C3 send the inputs, the model output and short passages of the source articles to {GENERATOR[0]} ({GENERATOR[1]}).")

    if st.button("\U0001f4ac Generate explanation", key="llm_generate"):
        try:
            with st.spinner("Generating the explanation…"):
                st.session_state["explanation"] = get_explanation(values, mode)
        except Exception as exc:
            st.error(f"The explanation could not be generated: {exc}")
    result = st.session_state.get("explanation")
    if result and (result["prediction"]["inputs"] != predict({k: values[k] for k in INPUT_COLS})["inputs"]
                   or result["prediction"]["anchored"] != (anchor_of(values) is not None)):
        result = None                       # explanation of other inputs (batch row changed): do not show it
    if result:
        show_explanation(result, "single")

    if batch_ok is not None:
        st.divider()
        n = min(len(batch_ok), MAX_BATCH_EXPLANATIONS)
        if len(batch_ok) > MAX_BATCH_EXPLANATIONS:
            st.caption(f"Explanations for all rows are limited to the first {MAX_BATCH_EXPLANATIONS} rows.")
        if st.button(f"\U0001f504 Generate explanations for {n} rows", key="llm_generate_all"):
            parts, progress = [], st.progress(0.0)
            for i in range(n):
                try:
                    parts.append(f"Row {i + 1}\n" + report_text(get_explanation(row_values(batch_ok, i), mode)))
                except Exception as exc:
                    parts.append(f"Row {i + 1}\nNot generated: {exc}\n")
                progress.progress((i + 1) / n)
            progress.empty()
            st.session_state["batch_report"] = "\n".join(parts)
        if st.session_state.get("batch_report"):
            st.download_button("\U0001f4be Reports for all rows (.txt)", st.session_state["batch_report"], "cu_mp_reports_all_rows.txt", "text/plain")

    st.divider()
    s = model_summary()
    st.caption(f"**Limits.** The explanation does not change any number: values come from the model in panel B. Context is taken from "
               f"the {s['n_studies']} laboratory studies of the dataset and the field studies only. Mechanisms are linked to polymers "
               "by keyword matching; two authors checked a sample, and 83–85% of the relations were judged correct (manuscript, "
               "Text S5). Citations name articles of the corpus; check the article before "
               "quoting a finding.")



# ── Panel E – explanations of the evaluation (Section 3.1.4) ─────────────────

EVAL_FILES = {"cases": "llm_eval_v2_cases_final.csv", "explanations": "llm_eval_v2_explanations_final.csv",
              "scores": "llm_eval_v2_scores_final.csv", "summary": "llm_eval_v2_summary_final.csv",
              "experts": "expert_eval_summary.csv"}
STRATA = {"river_level": "river concentration", "low_laboratory": "lowest 5% of laboratory Ce", "in_range": "within the laboratory range",
          "above_range": "above the laboratory range"}
CONDITIONS = {"C0": "C0 numbers only", "C1": "C1 rule-based", "C2": "C2 retrieval", "C3": "C3 retrieval + knowledge graph"}
JUDGE_NAMES = {"independent_provider": "Independent judge (Claude Haiku 4.5)", "same_model_as_generator": "Generator as judge (GPT-4o mini)"}
DIMENSIONS = {"accuracy": "Accuracy", "mechanistic_depth": "Mechanistic depth", "literature_grounding": "Literature grounding",
              "condition_specificity": "Condition specificity", "readability": "Readability", "limits_communicated": "Limits communicated"}


@st.cache_data
def evaluation_tables() -> dict[str, pd.DataFrame] | None:
    paths = {k: config.REPORTS_DIR / f for k, f in EVAL_FILES.items()}
    if not all(p.exists() for p in paths.values()):
        return None
    return {k: pd.read_csv(p) for k, p in paths.items()}


def yes_no(x: Any) -> str:
    return "yes" if bool(x) else "no"


def panel_examples() -> None:
    t = evaluation_tables()
    if t is None:
        st.info("ℹ️ The evaluation tables (outputs/reports/llm_eval_v2_*_final.csv) are not available.")
        return
    st.markdown("The 120 explanations evaluated in the manuscript (Section 3.1.4; Table 2; Fig. 4): 30 test cases, each explained in "
                "the four conditions. C2 and C3 were written by GPT-4o mini from passages of the 41 source articles; the passages "
                "themselves are not shown because the articles cannot be redistributed. Nothing here is generated anew.")
    st.caption("These explanations were generated with the evaluated version of the system. Their field-scale range is the 95% "
               "confidence interval of the mean field-to-laboratory ratio; the prediction panel of this application now gives the "
               "wider interval for a river not among the field studies.")

    with st.expander("\U0001f4ca Scores by condition", expanded=False):
        sc = t["scores"]
        overall = sc.pivot_table(index="condition", columns="judge", values="overall_weighted", aggfunc="mean")
        overall = overall.rename(columns=JUDGE_NAMES)
        ex = t["experts"]
        ex = ex[ex["rater"] == "mean_of_raters"].set_index("condition")["overall"].rename("Two experts (15 cases)")
        table = overall.join(ex).rename(index=CONDITIONS).round(2)
        st.dataframe(table, width="stretch")
        sm = t["summary"].set_index("condition")
        checks = pd.DataFrame({
            "Predicted values stated (%)": 100 * sm[["qe_stated", "interval_stated", "field_value_stated"]].min(axis=1),
            "Extrapolation warning when needed (%)": 100 * sm["range_warning_recall"],
            "Sparse-polymer warning when needed (%)": 100 * sm["sparse_warning_recall"],
            "Citations per explanation": sm["citations_per_item"],
            "Citations to corpus articles (%)": (100 * sm["citations_in_corpus"] / sm["citations_total"]).where(sm["citations_total"] > 0),
            "Mechanisms named": sm["mechanisms_named"],
            "Words": sm["words"],
        }).rename(index=CONDITIONS).round(1)
        st.dataframe(checks, width="stretch")
        st.caption("Overall score: weighted mean of five criteria (1–5). Checks without a judge as in Fig. 4c.")

    cases = t["cases"].copy()
    strata = st.multiselect("Concentration stratum", list(STRATA), default=list(STRATA), format_func=lambda k: STRATA[k],
                            key="ex_strata")
    cases = cases[cases["stratum"].isin(strata)].reset_index(drop=True)
    if cases.empty:
        st.info("No case in the selected strata.")
        return
    idx = st.selectbox("Test case", range(len(cases)), key="ex_case",
                       format_func=lambda i: (f"{cases.loc[i, 'case_id']} | {cases.loc[i, 'ReT']}, {cases.loc[i, 'AgS'].lower()}, "
                                              f"{cases.loc[i, 'AdC'].lower()} | Ce = {cases.loc[i, 'Ce']:.3g} mg/L | "
                                              f"{STRATA[cases.loc[i, 'stratum']]}"))
    case_id = cases.loc[idx, "case_id"]
    items = t["explanations"][t["explanations"]["case_id"] == case_id].set_index("condition")
    with st.expander("Model output given to every condition (facts)", expanded=False):
        st.text(items.iloc[0]["facts"])
    tabs = st.tabs([CONDITIONS[c] for c in CONDITIONS if c in items.index])
    for tab, cond in zip(tabs, [c for c in CONDITIONS if c in items.index]):
        row = items.loc[cond]
        with tab:
            text = str(row["explanation"])
            if cond == "C0":
                st.code(text, language=None)
            else:
                st.html(safe_html(text))
            if isinstance(row.get("sources"), str) and row["sources"].strip():
                labels = source_labels([{"paper_id": k.strip()} for k in row["sources"].split(";") if k.strip()])
                with st.expander(f"\U0001f4da Articles used as context ({len(labels)})"):
                    for label in labels:
                        st.markdown(f"- {label}")
            checks = {
                "Predicted values stated": yes_no(row["qe_stated"] and row["interval_stated"] and row["field_value_stated"]),
                "Extrapolation warning": ("given" if row["range_warning_given"] else "missing") if row["range_warning_needed"] else "not needed",
                "Sparse-polymer warning": ("given" if row["sparse_warning_given"] else "missing") if row["sparse_warning_needed"] else "not needed",
                "Values not matching the model": int(row["n_unsupported_qe_values"]),
                "Citations (in corpus / naming the polymer)": f"{int(row['n_citations'])} ({int(row['n_citations_in_corpus'])} / {int(row['n_citations_naming_polymer'])})",
                "Regenerated after the numerical check": yes_no(row["regenerated_after_check"]),
                "Words": int(row["n_words"]),
            }
            st.dataframe(pd.DataFrame([(k, str(v)) for k, v in checks.items()], columns=["Check without a judge", "Result"]),
                         hide_index=True, width="stretch")
            sc = t["scores"][t["scores"]["item_id"] == row["item_id"]]
            if not sc.empty:
                judged = sc.set_index("judge")[list(DIMENSIONS) + ["overall_weighted"]].rename(
                    index=JUDGE_NAMES, columns={**DIMENSIONS, "overall_weighted": "Overall"})
                st.dataframe(judged, width="stretch")
                with st.expander("Rationale of the judges"):
                    for _, r in sc.iterrows():
                        st.markdown(f"**{JUDGE_NAMES.get(r['judge'], r['judge'])}:** {r['rationale']}")


def main() -> None:
    st.title("Cu²⁺ – Microplastic Adsorption Predictor")
    st.caption("Prediction with study-level uncertainty · Contributions of the inputs · Field-scale value · LLM interpretation (C1/C2/C3) · Evaluated explanations")
    values, batch = render_sidebar()
    with st.expander("\U0001f4cb Panel A — Input summary", expanded=True):
        panel_input_summary(values, batch)
    st.divider()
    tab_b, tab_c, tab_d, tab_e = st.tabs(["\U0001f4ca B — Prediction", "\U0001f9e0 C — Contributions", "\U0001f4ac D — LLM",
                                          "\U0001f4d1 E — Evaluated explanations"])
    with tab_b:
        if batch:
            st.session_state.pop("prediction", None)
        panel_prediction(values, batch)
    with tab_c:
        panel_drivers()
    with tab_d:
        panel_llm()
    with tab_e:
        panel_examples()
    st.caption("The evaluation of the explanation module is reported in the manuscript and its Supplementary Information.")


if __name__ == "__main__":
    main()
