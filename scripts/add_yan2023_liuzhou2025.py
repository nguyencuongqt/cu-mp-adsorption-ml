"""Append Yan et al. (2023) and Liu Y et al. (2025) + Zhou W et al. (2025) to data/raw_data_literature_master.xlsx.

Source: Web of Science searches of 14 October 2026 (main string and supplementary S2a/S2b); eligibility in
Revision/05_Literature_search/Wos/Full_text_eligibility_WoS_S2_2026-10-14.md.

- Yan_2023 (10.1007/s11356-023-27906-1): Fig. 4, qe vs initial concentration, digitized by the authors (route A);
  Ce = C0 - 10 g/L x qe; 32 of 36 points after QC (C0 = 1-2 mg/L dropped for PLA and PE). Virgin PE, PP, PVC, PLA,
  d50 125 um; Milli-Q; pH 5.5; 296 K; 160 rpm; 48 h; single-metal solutions.
- LiuZhou_2025: Liu Y et al. 2025 (Land Degrad. Dev., 10.1002/ldr.5392; Fig. 6c, 7c: CK, thermal, UV) and Zhou W et al. 2025
  (J. Environ. Chem. Eng., 10.1016/j.jece.2025.120355; Fig. 6c, 7c: acid). Same laboratory and protocol; the pristine
  isotherms are shared (identical Table 5), so they are entered once (from Liu) and both reports form one study.
  x-axis labelled Ce but markers at nominal C0; Ce = C0 - 1 g/L x qe; 48 of 48 points. PS and PLA, 150 um; ultrapure
  water; 25 C; 160 rpm; 8 h; pH not stated.

QC files: Revision/06_Data_audit/Digitization_WoS_S2_new_papers/1[45]_*/rows_for_master_*.csv and QC_note_2026-10-14.md.
A backup of the data files is written first. Run scripts/build_grouped_dataset.py afterwards.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
MASTER = DATA / "raw_data_literature_master.xlsx"
BACKUP = DATA / "_backups" / "backup_2026-10-14_before_adding_Yan2023_LiuZhou2025"
DIG = ROOT / "Manuscript" / "Submission_Env pollution" / "Revision" / "06_Data_audit" / "Digitization_WoS_S2_new_papers"
FILES = {
    "Yan_2023": DIG / "14_ESPR_2023_27906_Yan_PE_PP_PVC_PLA" / "rows_for_master_14_ESPR_2023_27906_Yan_PE_PP_PVC_PLA.csv",
    "LiuZhou_2025": DIG / "15_LDD_2025_5392_and_JECE_2025_120355_LiuZhou_PS_PLA"
    / "rows_for_master_15_LDD_2025_5392_and_JECE_2025_120355_LiuZhou_PS_PLA.csv",
}
STUDIES = {
    "Yan_2023": {"ref": "(Yan et al., 2023)", "author": "Yan", "year": 2023, "doi": "10.1007/s11356-023-27906-1",
                 "paper_id": "Yan_2023_PE_PP_PVC_PLA_six_metals_MilliQ"},
    "LiuZhou_2025": {"ref": "(Liu et al., 2025; Zhou et al., 2025)", "author": "Liu", "year": 2025,
                     "doi": "10.1002/ldr.5392; 10.1016/j.jece.2025.120355",
                     "paper_id": "LiuY_2025_PS_PLA_thermal_UV_aging_CdCuZn; ZhouW_2025_PS_PLA_acid_aging_CdCuZn"},
}
STATUS = "verified_original_pdf_figure_and_qc"
AGEING = {  # template series_id -> (series_id in master, label, ageing_type, ageing_detail, AgS)
    "PLA_virgin": ("Yan_2023_PLA_virgin", "PLA 125um | virgin", "none", None, "Virgin"),
    "PE_virgin": ("Yan_2023_PE_virgin", "PE 125um | virgin", "none", None, "Virgin"),
    "PP_virgin": ("Yan_2023_PP_virgin", "PP 125um | virgin", "none", None, "Virgin"),
    "PVC_virgin": ("Yan_2023_PVC_virgin", "PVC 125um | virgin", "none", None, "Virgin"),
    "PS_virgin": ("LiuZhou_2025_PS_virgin", "PS 150um | virgin (CK)", "none", None, "Virgin"),
    "PLA_virgin_LZ": ("LiuZhou_2025_PLA_virgin", "PLA 150um | virgin (CK)", "none", None, "Virgin"),
    "PS_thermal30d": ("LiuZhou_2025_PS_thermal", "PS 150um | thermal 40 C/RT 30 d", "thermal", "12 h 40 C / 12 h room temperature cycles, 30 d", "Aged"),
    "PLA_thermal30d": ("LiuZhou_2025_PLA_thermal", "PLA 150um | thermal 40 C/RT 30 d", "thermal", "12 h 40 C / 12 h room temperature cycles, 30 d", "Aged"),
    "PS_UV30d": ("LiuZhou_2025_PS_UV", "PS 150um | UV 30 d", "UV", "UVA-340, 12 h light / 12 h dark, 25 C, 30 d", "Aged"),
    "PLA_UV30d": ("LiuZhou_2025_PLA_UV", "PLA 150um | UV 30 d", "UV", "UVA-340, 12 h light / 12 h dark, 25 C, 30 d", "Aged"),
    "PS_acid30d": ("LiuZhou_2025_PS_acid", "PS 150um | acid HNO3 pH 5.6 30 d", "acid", "dilute HNO3 pH 5.6, dark, solution renewed every 2 d, 30 d", "Aged"),
    "PLA_acid30d": ("LiuZhou_2025_PLA_acid", "PLA 150um | acid HNO3 pH 5.6 30 d", "acid", "dilute HNO3 pH 5.6, dark, solution renewed every 2 d, 30 d", "Aged"),
}


def new_rows() -> pd.DataFrame:
    out = []
    for sid, path in FILES.items():
        d = pd.read_csv(path, encoding="utf-8-sig")
        for _, r in d.iterrows():
            key = r["series_id"]
            if sid == "LiuZhou_2025" and key == "PLA_virgin":
                key = "PLA_virgin_LZ"
            series_id, label, atype, detail, ags = AGEING[key]
            out.append({
                "source_id": sid, "doi": r["doi"], "AgS": ags, "ReT": r["polymer"], "AdC": "Single-ion",
                "Temp": 23.0 if sid == "Yan_2023" else 25.0,
                "pH": 5.5 if sid == "Yan_2023" else None, "rpm": 160.0,
                "Ce": round(float(r["ce_mg_L"]), 6), "qe": round(float(r["qe_mg_g"]), 6),
                "series_id": series_id, "series_label": label, "ageing_type": atype, "ageing_detail": detail,
            })
    rows = pd.DataFrame(out)
    assert (rows["Ce"] > 0).all() and (rows["qe"] > 0).all()
    return rows


def main() -> None:
    wb = load_workbook(MASTER)
    ws = wb["qe_model_cu"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    existing = {ws.cell(row=r, column=col["source_id"]).value for r in range(2, ws.max_row + 1)}
    series_before = {ws.cell(row=r, column=col["series_id"]).value for r in range(2, ws.max_row + 1)}
    if existing & set(STUDIES):
        raise SystemExit("These studies are already in the master workbook.")
    rows = new_rows()
    clash = set(rows["series_id"]) & series_before
    assert not clash, clash

    BACKUP.mkdir(parents=True, exist_ok=True)
    for name in ["raw_data_literature_master.xlsx", "dataset.csv", "dataset_grouped.csv", "paper_metadata.csv", "size_ssa_metadata.csv"]:
        shutil.copy2(DATA / name, BACKUP / name)

    n_before = ws.max_row - 1
    next_id = max(ws.cell(row=r, column=col["raw_row_id"]).value for r in range(2, ws.max_row + 1)) + 1
    for i, r in rows.iterrows():
        s = STUDIES[r["source_id"]]
        record = {"source_id": r["source_id"], "source_author_year": s["ref"], "first_author": s["author"], "year": s["year"],
                  "doi": r["doi"], "doi_match_status": STATUS, "raw_row_id": next_id + i, "model_include": True,
                  "exclusion_reason": None, "AgS": r["AgS"], "ReT": r["ReT"], "AdC": r["AdC"], "Temp": r["Temp"],
                  "pH": r["pH"], "rpm": r["rpm"], "Ce": r["Ce"], "qe": r["qe"], "series_id": r["series_id"],
                  "series_label": r["series_label"], "ageing_type": r["ageing_type"], "ageing_detail": r["ageing_detail"]}
        ws.append([record[h] for h in header])

    meta = wb["literature_metadata"]
    mh = [c.value for c in meta[1]]
    for sid, s in STUDIES.items():
        part = rows[rows["source_id"] == sid]
        rec = {"source_id": sid, "source_author_year": s["ref"], "first_author": s["author"], "year": s["year"],
               "matched_paper_id": s["paper_id"], "doi": s["doi"], "n_cu_rows": len(part), "n_model_rows_included": len(part),
               "metals": "Cu (II)", "polymers": "; ".join(sorted(part["ReT"].unique()))}
        meta.append([rec[h] for h in mh])

    n_after, n_studies = ws.max_row - 1, len(existing) + len(STUDIES)
    n_series_before, n_series = len(series_before), len(series_before | set(rows["series_id"]))
    swaps = [(f"{n_before} Cu rows", f"{n_after} Cu rows"), (f"same {n_before} rows", f"same {n_after} rows"),
             (f"{len(existing)} literature studies", f"{n_studies} literature studies"),
             (f"({n_series_before} isotherm series)", f"({n_series} isotherm series)"),
             (f"All {len(existing)} studies", f"All {n_studies} studies"),
             ("Data collection closed 2026-10-07 (last studies added: Godoy_2019, Liu_2022; Mao_2020 removed).",
              "Data collection closed 2026-10-14 (last studies added: Yan_2023, LiuZhou_2025 from the Web of Science searches).")]
    for sheet in ("README", "validation_summary"):
        for row in wb[sheet].iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    for old, new in swaps:
                        cell.value = cell.value.replace(old, new)
                elif cell.value == n_before and sheet == "validation_summary":
                    cell.value = n_after
                elif cell.value == len(existing) and sheet == "validation_summary":
                    cell.value = n_studies
    wb.save(MASTER)

    pm = pd.read_csv(DATA / "paper_metadata.csv")
    add = pd.DataFrame([{"paper_id": "Yan_2023_PE_PP_PVC_PLA_six_metals_MilliQ", "doi": "10.1007/s11356-023-27906-1"},
                        {"paper_id": "LiuY_2025_PS_PLA_thermal_UV_aging_CdCuZn", "doi": "10.1002/ldr.5392"},
                        {"paper_id": "ZhouW_2025_PS_PLA_acid_aging_CdCuZn", "doi": "10.1016/j.jece.2025.120355"}])
    pd.concat([pm, add], ignore_index=True).to_csv(DATA / "paper_metadata.csv", index=False)
    print(f"rows {n_before} -> {n_after}; studies {len(existing)} -> {n_studies}; series {n_series_before} -> {n_series}; "
          f"raw_row_id {next_id}-{next_id + len(rows) - 1}")
    print(rows.groupby(["source_id", "series_id"]).agg(n=("Ce", "size"), Ce_min=("Ce", "min"), Ce_max=("Ce", "max"),
                                                        qe_min=("qe", "min"), qe_max=("qe", "max")).round(4).to_string())


if __name__ == "__main__":
    main()
