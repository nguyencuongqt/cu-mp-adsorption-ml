"""Append Turner & Holmes (2015), Environ. Chem., doi:10.1071/EN14143 to data/raw_data_literature_master.xlsx.

Source: field-study search of 16 October 2026 (Scopus/WoS, string v3), record #40; eligibility in
Revision/06_Data_audit/Digitization_field_search_2026-10/16_EnvChem_2015_Turner_Holmes_PE_pellets/Eligibility_review_Turner_2015.md.

Fig. 3 (virgin PE pellets) and Fig. 4 (beached PE pellets): adsorbed Cu ([MeX]*, ug/g; 1.2 M HCl extract) against measured
aqueous Cu ([Me], ug/L) after 48 h in filtered Plym river water (pH 6.3-6.5), 20 pellets in 50 mL, 19 +/- 2 C, 150 rpm,
nine metals added together at 2-20 ug/L. Digitized by the authors (17 Oct 2026); the beached-pellet readings were entered
with the y axis calibrated at 0.08 instead of 0.80 and were multiplied by 10 (file note); after correction all 10 points
agree with the authors' Freundlich fits (ratio 0.86-1.26). QC: scripts/check_turner2015_digitized.py.

A backup of the data files is written first. Run scripts/build_grouped_dataset.py and scripts/build_size_ssa_metadata.py afterwards.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
MASTER = DATA / "raw_data_literature_master.xlsx"
BACKUP = DATA / "_backups" / "backup_2026-10-17_before_adding_Turner2015"
DIG = (ROOT / "Manuscript" / "Submission_Env pollution" / "Revision" / "06_Data_audit" / "Digitization_field_search_2026-10"
       / "16_EnvChem_2015_Turner_Holmes_PE_pellets")
SID = "Turner_2015"
REF = "(Turner and Holmes, 2015)"
DOI = "10.1071/EN14143"
SERIES = {  # curve_id in rows_for_master -> (series_id, label, ageing_type, ageing_detail, AgS)
    "Turner_2015_01": ("Turner_2015_PE_virgin", "PE pellets ~4 mm | virgin", "none", None, "Virgin"),
    "Turner_2015_02": ("Turner_2015_PE_beached", "PE pellets 2-6 mm | beached (photo-oxidised)", "outdoor_weathering",
                       "pre-production pellets collected from the strandline of Watergate Bay (UK); carbonyl groups, yellowing",
                       "Aged"),
}


def main() -> None:
    rows = pd.read_csv(DIG / "rows_for_master_Turner_2015.csv")
    assert len(rows) == 10 and (rows["Ce"] > 0).all() and (rows["qe"] > 0).all()
    wb = load_workbook(MASTER)
    ws = wb["qe_model_cu"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    existing = {ws.cell(row=r, column=col["source_id"]).value for r in range(2, ws.max_row + 1)}
    series_before = {ws.cell(row=r, column=col["series_id"]).value for r in range(2, ws.max_row + 1)}
    if SID in existing:
        raise SystemExit("Turner_2015 is already in the master workbook.")

    BACKUP.mkdir(parents=True, exist_ok=True)
    for name in ["raw_data_literature_master.xlsx", "dataset.csv", "dataset_grouped.csv", "paper_metadata.csv",
                 "size_ssa_metadata.csv", "field_data.csv", "field_evidence_tiers.csv"]:
        shutil.copy2(DATA / name, BACKUP / name)

    n_before = ws.max_row - 1
    next_id = max(ws.cell(row=r, column=col["raw_row_id"]).value for r in range(2, ws.max_row + 1)) + 1
    for i, r in rows.reset_index(drop=True).iterrows():
        series_id, label, atype, detail, ags = SERIES[r["curve_id"]]
        record = {"source_id": SID, "source_author_year": REF, "first_author": "Turner", "year": 2015, "doi": DOI,
                  "doi_match_status": "verified_original_pdf_figure_and_qc", "raw_row_id": next_id + i, "model_include": True,
                  "exclusion_reason": None, "AgS": ags, "ReT": "PE", "AdC": "Mixed", "Temp": 19.0, "pH": 6.4, "rpm": 150.0,
                  "Ce": round(float(r["Ce"]), 7), "qe": round(float(r["qe"]), 8), "series_id": series_id,
                  "series_label": label, "ageing_type": atype, "ageing_detail": detail}
        ws.append([record[h] for h in header])

    meta = wb["literature_metadata"]
    mh = [c.value for c in meta[1]]
    rec = {"source_id": SID, "source_author_year": REF, "first_author": "Turner", "year": 2015,
           "matched_paper_id": "Turner_2015_PE_pellets_river_water_nine_metals", "doi": DOI, "n_cu_rows": 10,
           "n_model_rows_included": 10, "metals": "Cu (II) (with Ag, Cd, Co, Cr, Hg, Ni, Pb, Zn)", "polymers": "PE"}
    meta.append([rec.get(h) for h in mh])

    n_after, n_studies = ws.max_row - 1, len(existing) + 1
    n_series_before, n_series = len(series_before), len(series_before) + 2
    swaps = [(f"{n_before} Cu rows", f"{n_after} Cu rows"), (f"same {n_before} rows", f"same {n_after} rows"),
             (f"{len(existing)} literature studies", f"{n_studies} literature studies"),
             (f"({n_series_before} isotherm series)", f"({n_series} isotherm series)"),
             (f"All {len(existing)} studies", f"All {n_studies} studies"),
             ("Data collection closed 2026-10-14 (last studies added: Yan_2023, LiuZhou_2025 from the Web of Science searches).",
              "Data collection closed 2026-10-17 (last study added: Turner_2015, from the field-study search of 16 Oct 2026).")]
    for sheet in ("README", "validation_summary"):
        for row in wb[sheet].iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    for old, new in swaps:
                        cell.value = cell.value.replace(old, new)
                elif cell.value == n_before and sheet == "validation_summary":
                    cell.value = n_after
                elif cell.value == len(existing) and sheet == "validation_summary":
                    cell.value = n_studies
    wb.save(MASTER)

    pm = pd.read_csv(DATA / "paper_metadata.csv")
    add = pd.DataFrame([{"paper_id": "Turner_2015_PE_pellets_river_water_nine_metals", "doi": DOI}])
    pd.concat([pm, add], ignore_index=True).to_csv(DATA / "paper_metadata.csv", index=False)
    print(f"rows {n_before} -> {n_after}; studies {len(existing)} -> {n_studies}; series {n_series_before} -> {n_series}")


if __name__ == "__main__":
    main()
