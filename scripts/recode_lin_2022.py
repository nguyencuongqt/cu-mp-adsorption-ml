"""Recode Lin_2022 as a single-ion system and add its pH (decision of the authors, 2026-10-09).

Source: Supplementary Text S2 of Lin et al. (2022, Water Res. 225, 119144). Stock solutions of Cu(II), Ni(II) and Zn(II)
were prepared separately from the three nitrates; isotherms: 0.3 g of particles in 100 mL (3 g/L), seven concentrations
(15-350 umol/L), pH adjusted to 5.0, 170 rpm, 288-308 K. Neither the article nor the SI describes a solution containing
the three metals together, and competition between them is not discussed, so the 21 observations are coded Single-ion
(they were coded Mixed). The particles are polystyrene containing 2% HBCD as an additive (incorporated in production).
A backup of the data files is written first. Run scripts/build_grouped_dataset.py afterwards.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
MASTER = DATA / "raw_data_literature_master.xlsx"
BACKUP = DATA / "_backups" / "backup_2026-10-09_before_Lin2022_recode"
STUDY = "Lin_2022"


def main() -> None:
    BACKUP.mkdir(exist_ok=True)
    for name in ["raw_data_literature_master.xlsx", "dataset.csv", "dataset_grouped.csv"]:
        shutil.copy2(DATA / name, BACKUP / name)
    wb = load_workbook(MASTER)
    ws = wb["qe_model_cu"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    n = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=col["source_id"]).value != STUDY:
            continue
        if ws.cell(row=r, column=col["AdC"]).value != "Mixed":
            raise SystemExit(f"row {r}: AdC is not Mixed; nothing changed.")
        ws.cell(row=r, column=col["AdC"]).value = "Single-ion"
        ws.cell(row=r, column=col["pH"]).value = 5.0
        label = ws.cell(row=r, column=col["series_label"]).value
        ws.cell(row=r, column=col["series_label"]).value = label.replace(" | mixed", "")
        ws.cell(row=r, column=col["ageing_detail"]).value = "PS containing 2% HBCD as an additive (not aged)"
        n += 1
    wb.save(MASTER)
    print(f"{STUDY}: {n} rows recoded to Single-ion, pH 5.0; backup: {BACKUP.name}")


if __name__ == "__main__":
    main()
